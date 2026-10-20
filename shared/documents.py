"""Document creation tools (Word, PowerPoint, PDF, Excel) that run inside the agent container.

Each tool builds the file from structured input, stores it (blob SAS link, see generated_files.store_file)
and returns a markdown download link for the model to include in its reply.

Section shape (Word/PDF/PowerPoint): {"heading": str, "text": str, "bullets": [str], "table": [[str, ...], ...]}
where the first table row is the header.
"""

from __future__ import annotations

import io
from typing import Annotated, Any

from agent_framework import tool

from shared.generated_files import safe_name, store_file

Sections = Annotated[
    list[dict[str, Any]],
    'List of sections: {"heading": str, "text": str, "bullets": [str], "table": [[header...], [row...]]}. All keys optional.',
]


def _rows(table: Any) -> list[list[str]]:
    return [[str(c) for c in row] for row in table or [] if isinstance(row, (list, tuple))]


def _bullets(section: dict) -> list[str]:
    return [str(b) for b in section.get("bullets") or []]


def build_word(title: str, sections: list[dict]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    for s in sections:
        if s.get("heading"):
            doc.add_heading(str(s["heading"]), 1)
        if s.get("text"):
            doc.add_paragraph(str(s["text"]))
        for b in _bullets(s):
            doc.add_paragraph(b, style="List Bullet")
        rows = _rows(s.get("table"))
        if rows:
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows), style="Table Grid")
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_powerpoint(title: str, slides: list[dict]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    for s in slides:
        rows = _rows(s.get("table"))
        slide = prs.slides.add_slide(prs.slide_layouts[5 if rows else 1])
        slide.shapes.title.text = str(s.get("heading") or "")
        lines = ([str(s["text"])] if s.get("text") else []) + _bullets(s)
        if rows:
            width = max(len(r) for r in rows)
            shape = slide.shapes.add_table(len(rows), width, Inches(0.5), Inches(1.8), Inches(9), Inches(0.4) * len(rows))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    shape.table.cell(i, j).text = cell
        elif lines:
            frame = slide.placeholders[1].text_frame
            frame.text = lines[0]
            for line in lines[1:]:
                frame.add_paragraph().text = line
            for p in frame.paragraphs:
                p.font.size = Pt(20)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_pdf(title: str, sections: list[dict]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    story: list = [Paragraph(escape(title), styles["Title"])]
    for s in sections:
        if s.get("heading"):
            story.append(Paragraph(escape(str(s["heading"])), styles["Heading2"]))
        if s.get("text"):
            story.append(Paragraph(escape(str(s["text"])), styles["BodyText"]))
        if _bullets(s):
            story.append(ListFlowable([Paragraph(escape(b), styles["BodyText"]) for b in _bullets(s)], bulletType="bullet"))
        rows = _rows(s.get("table"))
        if rows:
            table = Table(rows, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 8))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, title=title).build(story)
    return buf.getvalue()


def build_excel(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for i, s in enumerate(sheets, 1):
        name = "".join(c for c in str(s.get("name") or f"Sheet{i}") if c not in "[]:*?/\\")[:31] or f"Sheet{i}"
        ws = wb.create_sheet(name)
        for r, row in enumerate(s.get("rows") or [], 1):
            for c, value in enumerate(row, 1):
                ws.cell(r, c, value)
        for cell in ws[1] if ws.max_row else []:
            cell.font = Font(bold=True)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _name(filename: str, ext: str) -> str:
    base = safe_name(filename, "document")
    return base if base.lower().endswith(ext) else base + ext


async def _finish(builder, filename: str, ext: str, *args) -> str:
    try:
        data = builder(*args)
    except Exception as exc:
        return f"Could not create the document: {type(exc).__name__}: {exc}"
    return await store_file(data, _name(filename, ext))


@tool(description="Create a Word (.docx) document and return a markdown download link to include in the reply.")
async def create_word(title: str, sections: Sections, filename: str = "report.docx") -> str:
    return await _finish(build_word, filename, ".docx", title, sections)


@tool(description="Create a PowerPoint (.pptx) deck (title slide plus one slide per section) and return a markdown download link.")
async def create_powerpoint(title: str, slides: Sections, filename: str = "presentation.pptx") -> str:
    return await _finish(build_powerpoint, filename, ".pptx", title, slides)


@tool(description="Create a PDF document and return a markdown download link to include in the reply.")
async def create_pdf(title: str, sections: Sections, filename: str = "report.pdf") -> str:
    return await _finish(build_pdf, filename, ".pdf", title, sections)


@tool(description='Create an Excel (.xlsx) workbook from sheets [{"name": str, "rows": [[cell,...],...]}] (first row is the header) and return a markdown download link.')
async def create_excel(
    sheets: Annotated[list[dict[str, Any]], 'List of {"name": str, "rows": [[...], ...]}'],
    filename: str = "workbook.xlsx",
) -> str:
    return await _finish(build_excel, filename, ".xlsx", sheets)


DOCUMENT_TOOLS = [create_word, create_powerpoint, create_pdf, create_excel]
